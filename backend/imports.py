"""Excel-first import preview. Nothing mutates until a validated preview is applied."""

import csv
import io
import json
import zipfile
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import ValidationError
from .models import (
    Order,
    Resource,
    Material,
    Calendar,
    Factory,
    Routing,
    Product,
    Supplier,
    Auxiliary,
    Customer,
    Model,
    Window,
)


class MaintenanceImport(Model):
    id: str
    resource_id: str
    start: datetime
    end: datetime
    reason: str


class OperatorImport(Auxiliary):
    kind: str = "Operator"


class ToolImport(Auxiliary):
    kind: str = "Tool"


SCHEMAS = {
    "Orders": (
        Order,
        [
            "id",
            "customer_id",
            "product_id",
            "quantity",
            "order_date",
            "requested_date",
            "priority",
        ],
    ),
    "Machines": (
        Resource,
        [
            "id",
            "name",
            "type",
            "department",
            "capacity",
            "calendar_id",
            "efficiency",
            "cost_per_hour",
        ],
    ),
    "Materials": (
        Material,
        [
            "id",
            "description",
            "stock",
            "reserved",
            "incoming",
            "arrival",
            "supplier_id",
            "safety_stock",
        ],
    ),
    "Shifts": (Calendar, ["id", "name", "weekdays", "shifts", "holidays"]),
    "Routing": (
        None,
        [
            "routing_id",
            "routing_name",
            "id",
            "name",
            "sequence",
            "resource_id",
            "cycle_minutes",
            "setup_minutes",
            "external_lead_minutes",
            "transit_minutes",
            "unit_cost",
            "required_skill",
            "tool_id",
            "transfer_minutes",
            "queue_minutes",
            "batch_size",
        ],
    ),
}
TABLE = {
    "Orders": "orders",
    "Machines": "resources",
    "Materials": "materials",
    "Shifts": "calendars",
    "Routing": "routings",
}
SCHEMAS.update(
    {
        "Products": (Product, ["id", "name", "routing_id", "batch_size", "materials"]),
        "Suppliers": (Supplier, ["id", "name", "lead_time_days"]),
        "Customers": (Customer, ["id", "name", "category", "priority_weight"]),
        "Operators": (
            OperatorImport,
            ["id", "name", "skill", "calendar_id", "capacity", "eligible_resources"],
        ),
        "Tools": (ToolImport, ["id", "name", "calendar_id", "capacity"]),
        "Maintenance": (
            MaintenanceImport,
            ["id", "resource_id", "start", "end", "reason"],
        ),
    }
)
TABLE.update(
    Products="products",
    Suppliers="suppliers",
    Customers="customers",
    Operators="auxiliaries",
    Tools="auxiliaries",
    Maintenance="resources",
)
SCHEMAS["Resources"] = SCHEMAS["Machines"]
TABLE["Resources"] = "resources"


def safe_cell(v):
    if isinstance(v, str) and (
        v.lstrip().startswith(("=", "+", "-", "@")) or v.startswith(("\t", "\r", "\n"))
    ):
        return "'" + v
    return v


def template(kind, factory):
    if kind not in SCHEMAS:
        raise ValueError("Unknown template")
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    headers = SCHEMAS[kind][1]
    ws.append(headers)
    if kind == "Routing":
        for r in factory.routings:
            for op in r.operations:
                for alt in op.alternatives:
                    row = dict(
                        routing_id=r.id,
                        routing_name=r.name,
                        **op.model_dump(),
                        **alt.model_dump(),
                    )
                    ws.append([safe_cell(row.get(h, "")) for h in headers])
    elif kind == "Maintenance":
        for r in factory.resources:
            for i, w in enumerate(r.unavailable):
                ws.append(
                    [
                        f"{r.id}-{i+1}",
                        r.id,
                        w.start.isoformat(),
                        w.end.isoformat(),
                        safe_cell(w.reason),
                    ]
                )
    else:
        for row in getattr(factory, TABLE[kind])[:2]:
            if kind in ("Operators", "Tools") and row.kind != (
                "Operator" if kind == "Operators" else "Tool"
            ):
                continue
            data = row.model_dump(mode="json")
            ws.append(
                [
                    safe_cell(
                        json.dumps(data[h])
                        if isinstance(data.get(h), (dict, list))
                        else data.get(h)
                    )
                    for h in headers
                ]
            )
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="164C40")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            38, max(18, len(str(col[0].value)) + 4)
        )
    notes = wb.create_sheet("Instructions")
    for line in [
        "PromiseFlow AI · import template",
        "Replace example rows with your data. Keep column names unchanged.",
        "IDs link to existing master records. Preview before applying.",
        "All dates use plant-local time: YYYY-MM-DDTHH:MM:SS.",
        "Shifts use JSON arrays of minutes: [[480,960]]. Weekdays: [0,1,2,3,4,5].",
        "Routing rows sharing a routing ID are grouped. Repeat operation IDs for alternate resources.",
        "Import replaces matching IDs and adds new ones. It does not delete missing records.",
        "Formula cells are rejected. Maximum 2,000 rows / 5 MB.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def preview(kind, content, filename, factory, mapping=None):
    if kind not in SCHEMAS:
        raise ValueError("Unknown import type")
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Maximum upload size is 5 MB")
    if filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if any(
                    "vbaproject" in name.lower() or "externallinks/" in name.lower()
                    for name in z.namelist()
                ):
                    raise ValueError(
                        "Macros and external workbook links are not allowed"
                    )
                if sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond 30 MB")
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            if kind not in wb.sheetnames:
                wb.close()
                raise ValueError(f"Workbook must contain a sheet named {kind}")
            ws = wb[kind]
            rows = []
            for row in ws.iter_rows():
                if any(c.data_type == "f" for c in row):
                    raise ValueError("Formula cells are not allowed in imports")
                rows.append([c.value for c in row])
                if len(rows) > 2001:
                    raise ValueError("Maximum 2,000 data rows")
            wb.close()
        except (zipfile.BadZipFile, KeyError):
            raise ValueError("Invalid XLSX workbook")
    elif filename.lower().endswith(".csv"):
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    else:
        raise ValueError("Use .xlsx or UTF-8 .csv")
    if not rows:
        raise ValueError("The file is empty")
    if len(rows) > 2001:
        raise ValueError("Maximum 2,000 data rows")
    headers = [str(x or "").strip() for x in rows[0]]
    if mapping:
        if not isinstance(mapping, dict) or any(
            k not in headers or not isinstance(v, str) for k, v in mapping.items()
        ):
            raise ValueError(
                "Column mapping must map existing headers to canonical field names"
            )
        headers = [mapping.get(h, h) for h in headers]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Duplicate or blank column names")
    model, required = SCHEMAS[kind]
    missing = [h for h in required if h not in headers]
    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))
    parsed = []
    errors = []
    display = []
    for number, values in enumerate(rows[1:], 2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        if any(v is not None and str(v).strip() for v in values[len(headers) :]):
            errors.append(
                dict(
                    row=number,
                    column="row",
                    message="Data extends beyond the declared headers",
                )
            )
            continue
        data = {
            h: v
            for h, v in zip(headers, values)
            if v is not None and str(v).strip() != ""
        }
        display.append(dict(row=number, values={k: str(v) for k, v in data.items()}))
        try:
            for key in (
                "weekdays",
                "shifts",
                "holidays",
                "materials",
                "eligible_resources",
                "unavailable",
            ):
                if key in data:
                    data[key] = json.loads(data[key])
            if model:
                old = (
                    next(
                        (
                            r.model_dump(mode="json")
                            for r in getattr(factory, TABLE[kind])
                            if r.id == data.get("id")
                        ),
                        {},
                    )
                    if kind != "Maintenance"
                    else {}
                )
                if kind in ("Tools", "Operators"):
                    expected_kind = "Tool" if kind == "Tools" else "Operator"
                    if old and old["kind"] != expected_kind:
                        raise ValueError("Existing auxiliary has a different kind")
                    data["kind"] = expected_kind
                parsed.append(
                    (
                        number,
                        model.model_validate({**old, **data}).model_dump(mode="json"),
                    )
                )
            else:
                parsed.append((number, data))
        except (ValidationError, ValueError, TypeError) as exc:
            if isinstance(exc, ValidationError):
                errors.extend(
                    dict(
                        row=number,
                        column=".".join(map(str, e["loc"])),
                        message=e["msg"],
                    )
                    for e in exc.errors()
                )
            else:
                errors.append(dict(row=number, column="row", message=str(exc)))
    records = []
    if kind == "Routing":
        groups = {}
        from .models import Alternative, Operation

        for row, data in parsed:
            try:
                rid = data.pop("routing_id")
                rname = data.pop("routing_name")
                g = groups.setdefault(rid, dict(id=rid, name=rname, operations={}))
                opid = data["id"]
                alt = Alternative(
                    **{
                        key: data.pop(key)
                        for key in list(data)
                        if key in Alternative.model_fields
                    }
                )
                if opid in g["operations"]:
                    prior = {
                        k: v
                        for k, v in g["operations"][opid].items()
                        if k != "alternatives"
                    }
                    incoming = Operation(**data, alternatives=[alt]).model_dump(
                        exclude={"alternatives"}
                    )
                    if prior != incoming or g["name"] != rname:
                        raise ValueError(
                            "Repeated routing rows disagree on operation fields"
                        )
                    g["operations"][opid]["alternatives"].append(alt.model_dump())
                else:
                    g["operations"][opid] = Operation(
                        **data, alternatives=[alt]
                    ).model_dump()
            except (ValueError, KeyError, TypeError) as exc:
                errors.append(dict(row=row, column="routing", message=str(exc)))
        for g in groups.values():
            try:
                records.append(
                    Routing(
                        id=g["id"],
                        name=g["name"],
                        operations=list(g["operations"].values()),
                    ).model_dump(mode="json")
                )
            except ValueError as exc:
                errors.append(dict(row=0, column="routing", message=str(exc)))
    else:
        records = [d for n, d in parsed]
        seen = set()
        for n, d in parsed:
            if d["id"] in seen:
                errors.append(dict(row=n, column="id", message="Duplicate ID in file"))
            seen.add(d["id"])
    merged = factory.model_dump(mode="json")
    if kind == "Maintenance":
        ids = set()
        for number, row in parsed:
            try:
                if row["id"] in ids:
                    raise ValueError("Duplicate maintenance ID in file")
                ids.add(row["id"])
                resource = next(
                    (r for r in merged["resources"] if r["id"] == row["resource_id"]),
                    None,
                )
                if resource is None:
                    raise ValueError("Unknown resource")
                window = Window(
                    start=row["start"], end=row["end"], reason=row["reason"]
                ).model_dump(mode="json")
                if window not in resource["unavailable"]:
                    resource["unavailable"].append(window)
            except ValueError as exc:
                errors.append(dict(row=number, column="maintenance", message=str(exc)))
        return dict(
            kind=kind,
            rows=display,
            errors=errors,
            added=len(parsed),
            updated=0,
            valid=bool(display) and not errors,
            merged=merged,
        )
    existing = {r["id"]: r for r in merged[TABLE[kind]]}
    adds = sum(r["id"] not in existing for r in records)
    existing.update({r["id"]: r for r in records})
    merged[TABLE[kind]] = list(existing.values())
    if not errors:
        try:
            Factory.model_validate(merged)
        except ValidationError as exc:
            errors.extend(
                dict(row=0, column="references", message=e["msg"]) for e in exc.errors()
            )
    if not display:
        errors.append(dict(row=0, column="file", message="No data rows"))
    return dict(
        kind=kind,
        rows=display,
        errors=errors,
        added=adds,
        updated=len(records) - adds,
        valid=not errors,
        merged=merged,
    )
