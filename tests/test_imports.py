import io
from openpyxl import load_workbook, Workbook
import pytest
from backend.imports import template, preview


@pytest.mark.parametrize(
    "kind", ["Orders", "Machines", "Routing", "Materials", "Shifts"]
)
def test_templates_are_valid_imports(kind, factory):
    data = template(kind, factory)
    result = preview(kind, data, kind + ".xlsx", factory)
    assert result["valid"], result["errors"]
    assert result["rows"]


def test_negative_quantity_identifies_row_and_column(factory):
    data = template("Orders", factory)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Orders"]
    ws["D2"] = -1
    out = io.BytesIO()
    wb.save(out)
    result = preview("Orders", out.getvalue(), "Orders.xlsx", factory)
    assert not result["valid"]
    assert any(e["row"] == 2 and e["column"] == "quantity" for e in result["errors"])


def test_formula_cells_rejected(factory):
    wb = load_workbook(io.BytesIO(template("Orders", factory)))
    wb["Orders"]["D2"] = "=1+1"
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(ValueError, match="Formula"):
        preview("Orders", out.getvalue(), "orders.xlsx", factory)


def test_invalid_foreign_key_rejected(factory):
    data = template("Orders", factory)
    wb = load_workbook(io.BytesIO(data))
    wb["Orders"]["B2"] = "UNKNOWN"
    out = io.BytesIO()
    wb.save(out)
    result = preview("Orders", out.getvalue(), "orders.xlsx", factory)
    assert not result["valid"] and result["errors"][0]["column"] == "references"
